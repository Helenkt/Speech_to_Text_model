from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.database import database
from app.services.auth import AuthError, hash_password, public_user


MAX_ACCOUNTS = 5_000

FIELD_ALIASES = {
    "username": {"username", "user_name", "ten_dang_nhap", "tai_khoan"},
    "password": {"password", "mat_khau"},
    "full_name": {"full_name", "name", "ho_ten", "ten"},
    "role": {"role", "vai_tro", "loai_tai_khoan"},
    "email": {"email", "thu_dien_tu"},
    "student_id": {
        "student_id",
        "student_code",
        "candidate_id",
        "ma_sinh_vien",
        "ma_ung_vien",
    },
    "active": {"active", "enabled", "hoat_dong", "trang_thai"},
}

ROLE_ALIASES = {
    "teacher": "teacher",
    "giao_vien": "teacher",
    "mentor": "teacher",
    "gv": "teacher",
    "candidate": "candidate",
    "ung_vien": "candidate",
    "student": "candidate",
    "sinh_vien": "candidate",
    "uv": "candidate",
    "sv": "candidate",
}


class AccountImportError(ValueError):
    pass


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _slug(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _canonical_record(raw: dict[str, Any]) -> dict[str, Any]:
    alias_to_field = {
        alias: field for field, aliases in FIELD_ALIASES.items() for alias in aliases
    }
    result: dict[str, Any] = {}
    for key, value in raw.items():
        field = alias_to_field.get(_slug(key))
        if field and field not in result:
            result[field] = value
    return result


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _active(value: Any) -> bool:
    if value is None or _text(value) == "":
        return True
    if isinstance(value, bool):
        return value
    normalized = _slug(value)
    if normalized in {"1", "true", "yes", "active", "hoat_dong", "co"}:
        return True
    if normalized in {"0", "false", "no", "inactive", "ngung", "khong"}:
        return False
    raise AccountImportError("'active' chỉ nhận true/false, 1/0 hoặc có/không")


def normalize_account(raw: Any, *, row_number: int) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise AccountImportError(f"Dòng {row_number}: tài khoản phải là một object")
    values = _canonical_record(raw)
    username = _text(values.get("username"))
    password = _text(values.get("password"))
    full_name = _text(values.get("full_name"))
    role_key = _slug(values.get("role"))
    email = _text(values.get("email"))
    student_id = _text(values.get("student_id"))

    if not username:
        raise AccountImportError(f"Dòng {row_number}: thiếu username")
    if len(username) < 3 or len(username) > 64 or re.search(r"\s", username):
        raise AccountImportError(
            f"Dòng {row_number}: username phải dài 3-64 ký tự và không có khoảng trắng"
        )
    if not full_name:
        raise AccountImportError(f"Dòng {row_number}: thiếu full_name/ho_ten")
    if role_key not in ROLE_ALIASES:
        raise AccountImportError(
            f"Dòng {row_number}: role phải là teacher/giao_vien hoặc candidate/ung_vien"
        )
    if email and ("@" not in email or len(email) > 254):
        raise AccountImportError(f"Dòng {row_number}: email không hợp lệ")
    if password and len(password) < 6:
        raise AccountImportError(f"Dòng {row_number}: mật khẩu phải có ít nhất 6 ký tự")

    return {
        "username": username,
        "password": password,
        "full_name": full_name[:120],
        "role": ROLE_ALIASES[role_key],
        "email": email[:254],
        "student_id": student_id[:80],
        "active": _active(values.get("active")),
        "row_number": row_number,
    }


def _json_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        document = json.loads(content.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise AccountImportError("Tệp JSON phải dùng mã hóa UTF-8") from exc
    except json.JSONDecodeError as exc:
        raise AccountImportError(f"JSON không hợp lệ: dòng {exc.lineno}") from exc
    if isinstance(document, dict):
        document = document.get("users", document.get("accounts"))
    if not isinstance(document, list):
        raise AccountImportError("JSON phải là một danh sách hoặc object có trường 'users'")
    return document


def _csv_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AccountImportError("Tệp CSV phải dùng mã hóa UTF-8") from exc
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise AccountImportError("CSV không có dòng tiêu đề")
    return list(reader)


def _xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise AccountImportError("Không đọc được file Excel .xlsx") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = None
        for row in rows:
            candidate = [_text(value) for value in row]
            canonical = set(_canonical_record({name: name for name in candidate}))
            if {"username", "full_name", "role"} <= canonical:
                headers = candidate
                break
        if not headers:
            raise AccountImportError("Excel không có dòng tiêu đề")
        result: list[dict[str, Any]] = []
        for row in rows:
            if not any(_text(value) for value in row):
                if result:
                    break
                continue
            result.append(
                {
                    headers[index]: value
                    for index, value in enumerate(row)
                    if index < len(headers)
                }
            )
        return result
    finally:
        workbook.close()


def parse_accounts(content: bytes, *, filename: str) -> tuple[list[dict[str, Any]], str]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        rows = _json_rows(content)
    elif suffix == ".csv":
        rows = _csv_rows(content)
    elif suffix == ".xlsx":
        rows = _xlsx_rows(content)
    elif suffix == ".xls":
        raise AccountImportError("Hãy lưu file Excel cũ .xls thành định dạng .xlsx")
    else:
        raise AccountImportError("Chỉ hỗ trợ file .json, .csv hoặc .xlsx")

    if not rows:
        raise AccountImportError("Tệp không có tài khoản")
    if len(rows) > MAX_ACCOUNTS:
        raise AccountImportError(f"Mỗi lần chỉ nhập tối đa {MAX_ACCOUNTS} tài khoản")

    accounts: list[dict[str, Any]] = []
    usernames: set[str] = set()
    for index, row in enumerate(rows, start=2):
        account = normalize_account(row, row_number=index)
        key = account["username"].casefold()
        if key in usernames:
            raise AccountImportError(
                f"Dòng {index}: username '{account['username']}' bị trùng trong file"
            )
        usernames.add(key)
        accounts.append(account)
    return accounts, hashlib.sha256(content).hexdigest()


def import_accounts(
    database_path: Path,
    content: bytes,
    *,
    filename: str,
    mode: str = "upsert",
    imported_by: int | None = None,
) -> dict[str, Any]:
    if mode not in {"upsert", "replace"}:
        raise AccountImportError("Chế độ import phải là upsert hoặc replace")
    accounts, digest = parse_accounts(content, filename=filename)
    now = _now()
    created = 0
    updated = 0
    deactivated = 0

    with database(database_path) as connection:
        current = {
            row["username"].casefold(): row
            for row in connection.execute("SELECT * FROM users").fetchall()
        }
        for account in accounts:
            existing = current.get(account["username"].casefold())
            if not existing and not account["password"]:
                raise AccountImportError(
                    f"Dòng {account['row_number']}: tài khoản mới phải có password"
                )
            try:
                password_hash = (
                    hash_password(account["password"])
                    if account["password"]
                    else existing["password_hash"]
                )
            except AuthError as exc:
                raise AccountImportError(
                    f"Dòng {account['row_number']}: {exc}"
                ) from exc

            active = account["active"]
            if imported_by and existing and int(existing["user_id"]) == imported_by:
                active = True
            connection.execute(
                """
                INSERT INTO users (
                    username, password_hash, full_name, role, email, student_id,
                    active, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(username) DO UPDATE SET
                    password_hash = excluded.password_hash,
                    full_name = excluded.full_name,
                    role = excluded.role,
                    email = excluded.email,
                    student_id = excluded.student_id,
                    active = excluded.active,
                    updated_at = excluded.updated_at
                """,
                (
                    account["username"],
                    password_hash,
                    account["full_name"],
                    account["role"],
                    account["email"] or None,
                    account["student_id"] or None,
                    int(active),
                    now,
                    now,
                ),
            )
            if existing:
                updated += 1
            else:
                created += 1

        if mode == "replace":
            imported_names = {account["username"].casefold() for account in accounts}
            for row in connection.execute("SELECT user_id, username, active FROM users"):
                if row["username"].casefold() in imported_names:
                    continue
                if imported_by and int(row["user_id"]) == imported_by:
                    continue
                if row["active"]:
                    connection.execute(
                        "UPDATE users SET active = 0, updated_at = ? WHERE user_id = ?",
                        (now, row["user_id"]),
                    )
                    deactivated += 1

        teacher_count = connection.execute(
            "SELECT COUNT(*) AS total FROM users WHERE role = 'teacher' AND active = 1"
        ).fetchone()["total"]
        if not teacher_count:
            raise AccountImportError("Hệ thống phải còn ít nhất một Giáo viên hoạt động")

        connection.execute(
            """
            INSERT INTO account_sources (
                source_name, source_sha256, account_count, import_mode,
                imported_by, imported_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (filename, digest, len(accounts), mode, imported_by, now),
        )

    return {
        "imported": len(accounts),
        "created": created,
        "updated": updated,
        "deactivated": deactivated,
        "mode": mode,
        "source_sha256": digest,
    }


def list_accounts(database_path: Path) -> list[dict[str, Any]]:
    with database(database_path) as connection:
        rows = connection.execute(
            """
            SELECT * FROM users
            ORDER BY CASE role WHEN 'teacher' THEN 0 ELSE 1 END,
                     active DESC, full_name COLLATE NOCASE
            """
        ).fetchall()
    return [public_user(row) for row in rows]


def set_account_active(
    database_path: Path, user_id: int, *, active: bool, changed_by: int
) -> dict[str, Any]:
    if user_id == changed_by and not active:
        raise AccountImportError("Bạn không thể tự khóa tài khoản đang đăng nhập")
    with database(database_path) as connection:
        row = connection.execute(
            "SELECT * FROM users WHERE user_id = ?", (user_id,)
        ).fetchone()
        if not row:
            raise AccountImportError("Không tìm thấy tài khoản")
        if row["role"] == "teacher" and not active:
            teachers = connection.execute(
                "SELECT COUNT(*) AS total FROM users WHERE role = 'teacher' AND active = 1"
            ).fetchone()["total"]
            if int(teachers) <= 1:
                raise AccountImportError("Không thể khóa Giáo viên hoạt động cuối cùng")
        connection.execute(
            "UPDATE users SET active = ?, updated_at = ? WHERE user_id = ?",
            (int(active), _now(), user_id),
        )
        updated = connection.execute(
            "SELECT * FROM users WHERE user_id = ?", (user_id,)
        ).fetchone()
    return public_user(updated)


def seed_demo_accounts(database_path: Path, data_dir: Path) -> dict[str, Any]:
    path = data_dir / "accounts_demo.csv"
    if not path.exists():
        raise AccountImportError("Thiếu data/accounts_demo.csv")
    return import_accounts(
        database_path,
        path.read_bytes(),
        filename=path.name,
        mode="upsert",
    )
